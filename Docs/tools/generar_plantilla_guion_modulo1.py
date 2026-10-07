# -*- coding: utf-8 -*-
"""Genera la plantilla de guion del modulo 1 de RideSafe, con la estructura del
guion de Excavaciones y cada hueco atado a su clave real de I2."""
import io, sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from keys import KEYS

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

CORAL = RGBColor(0xE8, 0x4A, 0x33)
NAVY = RGBColor(0x14, 0x2A, 0x4F)
GREY = RGBColor(0x66, 0x6D, 0x75)

doc = Document()
for s in ("Normal",):
    f = doc.styles[s].font
    f.name = "Calibri"; f.size = Pt(11)


def h(text, level=1, color=NAVY):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = color
    return p


def para(text="", bold=False, italic=False, color=None, size=11, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color is not None: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def blank(label):
    """Hueco que rellena el cliente."""
    p = doc.add_paragraph()
    r = p.add_run("〔 " + label + " 〕")
    r.bold = True; r.font.color.rgb = CORAL
    p.paragraph_format.space_after = Pt(8)
    return p


def dialog(who, key=None):
    p = doc.add_paragraph()
    r = p.add_run("📜 " + who + ": ")
    r.bold = True
    if key:
        r2 = p.add_run("〔 POR DEFINIR 〕")
        r2.bold = True; r2.font.color.rgb = CORAL
        p.add_run("   ")
        r3 = p.add_run("clave: " + key)
        r3.font.size = Pt(8); r3.font.color.rgb = GREY
    else:
        r2 = p.add_run("〔 POR DEFINIR 〕")
        r2.bold = True; r2.font.color.rgb = CORAL
    return p


def keyline(key, nota=""):
    """Texto que YA existe en el build, con su valor actual."""
    en, es = KEYS.get(key, ("", ""))
    p = doc.add_paragraph()
    r = p.add_run("Ya existe · " + key + "  →  ")
    r.font.size = Pt(8); r.font.color.rgb = GREY
    r2 = p.add_run("“" + es + "”")
    r2.italic = True; r2.font.size = Pt(9)
    if nota:
        r3 = p.add_run("   " + nota)
        r3.font.size = Pt(8); r3.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(4)
    return p


def bullet(text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(text)
        r.bold = True; r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            if text.startswith("〔"):
                r = p.add_run(text); r.bold = True; r.font.color.rgb = CORAL; r.font.size = Pt(9)
            else:
                r = p.add_run(text); r.font.size = Pt(9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


# ───────────────────────── Portada ─────────────────────────
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("GUION\nREALIDAD VIRTUAL\nRIDESAFE — MOVILIDAD EN BICICLETA Y PATINETE ELÉCTRICO")
r.bold = True; r.font.size = Pt(22); r.font.color.rgb = NAVY

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MÓDULO 1 · PREPARACIÓN ANTES DE RODAR\nPLANTILLA PARA DILIGENCIAR")
r.bold = True; r.font.size = Pt(15); r.font.color.rgb = CORAL

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Versión 1 · 2026-10-07")
r.font.size = Pt(10); r.font.color.rgb = GREY

doc.add_page_break()

# ───────────────────── Cómo usar la plantilla ─────────────────────
h("CÓMO USAR ESTA PLANTILLA", 1)
para("Este documento sigue la misma estructura del guion de Excavaciones y Movimiento de "
     "Tierras, adaptada al módulo 1 de RideSafe, que ya está construido y funcionando.")
para("Hay dos tipos de bloque, y la diferencia importa:", bold=True)

bullet("〔 POR DEFINIR 〕 en coral — un hueco que usted escribe. Cada hueco indica la clave "
       "con la que ese texto entra al simulador.")
para("")
bullet("Ya existe · clave → “texto” en gris — ese texto YA está en el simulador, en inglés y "
       "español. Se puede cambiar: basta con escribir el nuevo valor al lado.")

para("")
para("Las claves marcadas como NUEVA no existen todavía: escribirlas implica desarrollo "
     "adicional. Están señaladas para que se pueda decidir el alcance con la información "
     "a la vista.", italic=True, color=GREY)

h("TRES COSAS QUE CONVIENE SABER ANTES DE ESCRIBIR", 2)
bullet("No hay avatar instructor. A diferencia del guion de Excavaciones, que se apoya en el "
       "avatar DANTE, hoy RideSafe usa narración en voz off con subtítulos palabra por palabra. "
       "Si se quiere un avatar visible, es desarrollo nuevo.")
bullet("No hay puntaje numérico. El módulo hoy clasifica cada elemento en cuatro categorías "
       "(Núcleo, Inapropiado, Opcional, Según condición) y cierra con una decisión go/no-go. "
       "Los rangos 100 / 75 / 50 que se piden al final requieren añadir un marcador numérico: "
       "es desarrollo nuevo, y por eso su apartado va marcado como NUEVA.")
bullet("Todo texto que escriba aquí se traduce a inglés y español. El simulador está en los dos "
       "idiomas y cambia en vivo.")

doc.add_page_break()

# ───────────────── Estructura narrativa ─────────────────
h("ESTRUCTURA NARRATIVA DEL GUION EN REALIDAD VIRTUAL", 1)
para("Apartados de referencia para el equipo. No forman parte de la escena.", italic=True, color=GREY)

h("INTRODUCCIÓN", 2)
para("Vista general del propósito formativo del módulo. No hace parte activa de la escena.",
     italic=True, color=GREY)
blank("POR DEFINIR — introducción del módulo 1")

h("OBJETIVO", 2)
para("Qué se busca evaluar, desarrollar o reforzar en el usuario.", italic=True, color=GREY)
blank("POR DEFINIR — objetivo del módulo 1")

h("CONTEXTO", 2)
para("Factores externos en juego y justificación de la escena. Base narrativa para el realismo.",
     italic=True, color=GREY)
blank("POR DEFINIR — contexto")

doc.add_page_break()

# ───────────────── Componentes funcionales ─────────────────
h("COMPONENTES FUNCIONALES DEL GUION", 1)
para("Elementos técnicos que dan forma a la escena. Los que ya están resueltos aparecen con su "
     "valor actual; los demás quedan abiertos.", italic=True, color=GREY)

table(
    ["Componente", "Estado en el simulador", "Qué debe definir el cliente"],
    [
        ["ESCENA", "Garaje doméstico cerrado, interior realista. El aprendiz entra desde fuera: "
                   "la puerta se abre y se le teletransporta dentro.", "〔 Ajustes si aplica 〕"],
        ["DURACIÓN", "3 a 4 minutos (según narrativa v2.0).", "〔 Confirmar 〕"],
        ["MODALIDAD", "Guiada. El aprendiz no se desplaza libremente: el módulo lo lleva de "
                      "estación en estación con fundido a negro.", "〔 Confirmar 〕"],
        ["INSTRUCTOR", "No hay avatar. Narración en voz off con subtítulos.",
         "〔 ¿Se quiere avatar? ¿Nombre y tono de voz? 〕"],
        ["UBICACIÓN", "Garaje de vivienda. El vehículo está dentro.", "〔 Confirmar 〕"],
        ["HORA DEL DÍA", "Día, luz clara. Es lo que hace que el chaleco reflectivo sea "
                         "“aceptable pero no exigido”.", "〔 Confirmar 〕"],
        ["ILUMINACIÓN", "Luz interior de garaje, cálida y uniforme.", "〔 Ajustes si aplica 〕"],
        ["SONIDO AMBIENTAL", "Pendiente de definir y producir.",
         "〔 POR DEFINIR — ambiente sonoro 〕"],
        ["DINÁMICA", "Dos entregas: (A) elegir el equipo de protección personal, "
                     "(B) inspeccionar el vehículo. Cada una cierra con su reporte.", "—"],
        ["DIÁLOGOS", "Voz off con subtítulos. Se marcan en cursiva en este documento.",
         "〔 POR DEFINIR — textos de locución 〕"],
    ],
    widths=[3.2, 7.0, 5.5])

doc.add_page_break()

# ───────────────── Contexto 1 ─────────────────
h("(CONTEXTO No 1)", 1)
for linea, valor in [
    ("Cliente / concesión", "〔 POR DEFINIR 〕"),
    ("Duración estimada", "3 a 4 min"),
    ("Escenario", "Garaje doméstico. El aprendiz entra por la puerta principal."),
    ("Hora del día", "Día, condiciones secas y despejadas"),
    ("Sonido ambiental", "〔 POR DEFINIR 〕"),
    ("Vista", "Modo PC con mouse hoy; Meta Quest como destino"),
    ("Vehículo", "Bicicleta eléctrica o patinete eléctrico, a elección del aprendiz en el módulo 0"),
    ("Avatar instructor", "〔 POR DEFINIR — hoy solo hay voz off 〕"),
]:
    p = doc.add_paragraph()
    r = p.add_run("🔹 " + linea + ": "); r.bold = True; r.font.size = Pt(10)
    if valor.startswith("〔"):
        r2 = p.add_run(valor); r2.bold = True; r2.font.color.rgb = CORAL; r2.font.size = Pt(10)
    else:
        r2 = p.add_run(valor); r2.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(3)

doc.add_page_break()

# ───────────────── Actividad ─────────────────
h("🔹 INICIO ACTIVIDAD 1 · MÓDULO 1", 1)

h("Pantalla de bienvenida", 2)
keyline("Module1/M01_Welcome", "(hoy sin usar en el módulo 1)")
dialog("Voz off — bienvenida", "NUEVA: Module1/Narr_Welcome")
para("Botón en pantalla: “Begin”. ", size=10)

h("Entrada al garaje", 2)
para("La puerta se abre, la pantalla funde a negro y el aprendiz aparece dentro, frente a la "
     "primera estación.", size=10)
dialog("Voz off — entrada", "NUEVA: Module1/Narr_Entrance")

h("Orientación", 2)
keyline("Module1/Orientation_Title")
keyline("Module1/Orientation_Body")
keyline("Module1/Start", "botón")
keyline("Module1/RepeatAudio", "botón")
dialog("Voz off — orientación", "NUEVA: Module1/Narr_Orientation")

doc.add_page_break()

# ───────────────── Sección A ─────────────────
h("SECCIÓN A · EQUIPO DE PROTECCIÓN PERSONAL", 1)
para("Tres estaciones. En cada una el aprendiz hace clic en un objeto, lo ve girar sobre un "
     "pedestal y decide si se lo lleva. Lo aceptado se acumula en la lista lateral y puede "
     "quitarse antes de enviar.", size=10)

keyline("Module1/Choose_Instruction", "instrucción permanente")
keyline("Module1/Confirm_Question", "pregunta del popup")
keyline("Module1/Yes")
keyline("Module1/No")

para("")
para("Para cada objeto hay que escribir el texto de voz off que explica POR QUÉ es adecuado o "
     "no. Hoy ese texto no existe: el aprendiz ve el objeto y el nombre, pero nadie le explica "
     "el criterio.", bold=True)

ITEMS = [
    ("ESTACIÓN 1 — CABEZA Y PROTECCIÓN", "Module1/Zone_Head", [
        ("Casco íntegro y ajustado", "Module1/Item_HelmetOk", "Núcleo", "✅ Llevar"),
        ("Casco deteriorado", "Module1/Item_HelmetDamaged", "Inapropiado", "❌ No llevar"),
        ("Gafas", "Module1/Item_Glasses", "Opcional", "Indiferente"),
    ]),
    ("ESTACIÓN 2 — ROPA Y CALZADO", "Module1/Zone_Clothing", [
        ("Calzado cerrado", "Module1/Item_ShoesOk", "Núcleo", "✅ Llevar"),
        ("Chaleco reflectivo", "Module1/Item_Reflective", "Según condición", "Opcional de día"),
    ]),
    ("ESTACIÓN 3 — CARGA Y MANOS LIBRES", "Module1/Zone_Load", [
        ("Carga asegurada", "Module1/Item_CargoSecured", "Núcleo", "✅ Llevar"),
        ("Objeto en la mano", "Module1/Item_CargoHandheld", "Inapropiado", "❌ No llevar"),
    ]),
]

for titulo, zonakey, items in ITEMS:
    h(titulo, 2)
    keyline(zonakey, "título de la estación")
    table(
        ["Elemento", "Categoría", "Respuesta esperada", "Voz off: por qué (NUEVA)"],
        [[n, cat, esp, "〔 POR DEFINIR 〕"] for (n, k, cat, esp) in items],
        widths=[4.5, 3.0, 3.2, 5.0])

para("Nota: la “ropa suelta o cordones sueltos” está prevista como riesgo de la estación 2, "
     "pero hoy está desactivada porque no hay modelo 3D. Mientras siga así, esa estación no "
     "tiene ninguna respuesta equivocada.", italic=True, color=GREY, size=9)

h("Lista lateral — “Mi preparación”", 2)
keyline("Module1/Sidebar_Title")
keyline("Module1/Sidebar_Empty")
keyline("Module1/Sidebar_Note")

h("Revisión antes de enviar", 2)
keyline("Module1/Review_Title")
keyline("Module1/Review_Subtitle")
keyline("Module1/Review_Lock")
keyline("Module1/Confirm", "botón")
keyline("Module1/KeepChoosing", "botón")

doc.add_page_break()

# ───────────────── Reporte ─────────────────
h("REPORTE DE LA SECCIÓN", 1)
para("Dos columnas: lo que eligió el aprendiz y la referencia para ese recorrido.", size=10)
keyline("Module1/Report_Title")
keyline("Module1/Report_Subtitle")
keyline("Module1/Report_Yours")
keyline("Module1/Report_Reference")
para("Leyenda:", bold=True, size=10)
for k in ("Module1/Selected", "Module1/Omitted", "Module1/Conditional", "Module1/Inappropriate"):
    keyline(k)
dialog("Voz off — retroalimentación de la sección", "NUEVA: Module1/Narr_ReportPersonal")

h("Video explicativo", 2)
keyline("Module1/WatchVideo", "botón")
keyline("Module1/VideoPlaceholder")
keyline("Module1/Pause")
keyline("Module1/Replay")
keyline("Module1/Continue")
blank("POR DEFINIR — guion del video de la sección de equipo personal")

doc.add_page_break()

# ───────────────── Sección B ─────────────────
h("SECCIÓN B · INSPECCIÓN DEL VEHÍCULO", 1)
para("Los mismos cinco chequeos para bicicleta y patinete: cambia dónde está cada pieza, no qué "
     "se revisa. El aprendiz marca cada punto del vehículo.", size=10)
keyline("Module1/Inspect_Instruction")
keyline("Module1/Inspect_Question")

PUNTOS = [
    ("ESTACIÓN 4 — DIRECCIÓN, FRENOS Y CONTROLES", "Module1/Zone_Cockpit", [
        ("Dirección y manubrio", "Module1/Point_Steering"),
        ("Frenos", "Module1/Point_Brakes"),
        ("Estado eléctrico y controles", "Module1/Point_Electrics"),
    ]),
    ("ESTACIÓN 5 — RUEDAS Y ESTRUCTURA", "Module1/Zone_Wheels", [
        ("Ruedas y estructura", "Module1/Point_Wheels"),
    ]),
    ("ESTACIÓN 6 — LUCES Y VISIBILIDAD", "Module1/Zone_Visibility", [
        ("Luces y reflectivos", "Module1/Point_Lights"),
    ]),
]

for titulo, zonakey, puntos in PUNTOS:
    h(titulo, 2)
    keyline(zonakey, "título de la estación")
    table(
        ["Punto a revisar", "Qué debe observar el aprendiz (NUEVA)", "Aviso si está mal (NUEVA)"],
        [[n, "〔 POR DEFINIR 〕", "〔 POR DEFINIR 〕"] for (n, k) in puntos],
        widths=[4.5, 5.6, 5.6])

h("Decisión final — ¿puede salir a rodar?", 2)
keyline("Module1/Decision_Question")
keyline("Module1/RoadReady", "botón")
keyline("Module1/NeedsService", "botón")
keyline("Module1/DoNotRide", "botón")
keyline("Module1/Decision_Correct", "si acierta")
keyline("Module1/Decision_Review", "si falla")
dialog("Voz off — cierre del módulo", "NUEVA: Module1/Narr_Closing")
keyline("Module1/Final_Complete")

doc.add_page_break()

# ───────────────── Puntaje ─────────────────
h("SISTEMA DE PUNTAJE", 1)
para("DESARROLLO NUEVO.", bold=True, color=CORAL)
para("Hoy el módulo no lleva marcador numérico: clasifica cada elemento en categorías y cierra "
     "con la decisión go/no-go. Para dar un porcentaje hay que definir cuánto suma o resta cada "
     "acción. La propuesta de abajo es un punto de partida; los valores quedan abiertos.", size=10)

table(
    ["Acción del aprendiz", "Puntos"],
    [
        ["Llevar un elemento núcleo (casco apto, calzado cerrado, carga asegurada)", "〔 POR DEFINIR 〕"],
        ["Omitir un elemento núcleo", "〔 POR DEFINIR 〕"],
        ["Llevar un elemento inapropiado (casco deteriorado, objeto en la mano)", "〔 POR DEFINIR 〕"],
        ["Rechazar correctamente un elemento inapropiado", "〔 POR DEFINIR 〕"],
        ["Elemento opcional (gafas): llevarlo o no", "〔 POR DEFINIR 〕"],
        ["Elemento según condición (chaleco) en día despejado", "〔 POR DEFINIR 〕"],
        ["Revisar un punto del vehículo", "〔 POR DEFINIR 〕"],
        ["Omitir un punto del vehículo", "〔 POR DEFINIR 〕"],
        ["Acertar la decisión final go/no-go", "〔 POR DEFINIR 〕"],
        ["Fallar la decisión final go/no-go", "〔 POR DEFINIR 〕"],
    ],
    widths=[11.0, 4.5])

h("RANGOS DE EVALUACIÓN", 1)
para("Cuatro mensajes de cierre, uno por franja. Son textos nuevos: cada uno necesita su clave.",
     size=10)

table(
    ["Franja", "Clave (NUEVA)", "Mensaje de cierre"],
    [
        ["100 %", "Module1/Score_Perfect", "〔 POR DEFINIR 〕"],
        ["75 % – 99 %", "Module1/Score_High", "〔 POR DEFINIR 〕"],
        ["50 % – 74 %", "Module1/Score_Mid", "〔 POR DEFINIR 〕"],
        ["Menos de 50 %", "Module1/Score_Low", "〔 POR DEFINIR 〕"],
    ],
    widths=[3.0, 4.5, 8.0])

para("Sugerencia de redacción, para mantener el tono del guion de Excavaciones: el mensaje de "
     "100 % reconoce el criterio; el de 75–99 % felicita pero señala el detalle omitido; el de "
     "50–74 % advierte sin castigar; el de menos de 50 % pide reforzar y repetir.",
     italic=True, color=GREY, size=9)

h("🔹 FIN ACTIVIDAD 1 · MÓDULO 1", 2)
bullet("“Repetir escenario”")
bullet("“Ver puntaje obtenido”")
bullet("“Avanzar al siguiente módulo”")
para("Corte en fondo negro.", italic=True, size=9, color=GREY)

doc.add_page_break()

# ───────────────── Anexo ─────────────────
h("ANEXO · INVENTARIO DE TEXTOS DEL MÓDULO 1", 1)
para("Los 61 textos que hoy existen en el simulador, con su valor actual. Para cambiar "
     "cualquiera basta con escribir el nuevo valor en la última columna.", size=10)

filas = [[k, es, "〔 　 〕"] for k, (en, es) in sorted(KEYS.items())]
table(["Clave", "Texto actual en español", "Nuevo texto"], filas, widths=[5.0, 7.0, 5.0])

out = "Docs/Guion_Modulo1_RideSafe_PLANTILLA.docx"
doc.save(out)
print("escrito:", out, os.path.getsize(out), "bytes")
